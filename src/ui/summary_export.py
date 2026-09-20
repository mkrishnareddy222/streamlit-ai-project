from io import BytesIO
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def create_pdf(summary: str, title: str = "Chat Summary") -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(title, styles["Title"]),
        Spacer(1, 0.2 * inch),
    ]
    for paragraph in summary.split("\n"):
        if paragraph.strip():
            story.append(Paragraph(escape(paragraph), styles["BodyText"]))
            story.append(Spacer(1, 0.1 * inch))
    document.build(story)
    return output.getvalue()


def create_pptx(summary: str, title: str = "Chat Summary") -> bytes:
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_slide.placeholders[1].text = "Generated from the chat conversation"

    paragraphs = [line.strip() for line in summary.split("\n") if line.strip()]
    for start in range(0, len(paragraphs), 6):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        text_box = slide.shapes.add_textbox(
            Inches(0.7), Inches(1.3), Inches(8.6), Inches(5.5)
        )
        frame = text_box.text_frame
        frame.word_wrap = True
        for index, paragraph in enumerate(paragraphs[start : start + 6]):
            if index == 0:
                text = frame.paragraphs[0]
            else:
                text = frame.add_paragraph()
            text.text = paragraph
            text.font.size = Pt(20)
            text.space_after = Pt(12)

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
